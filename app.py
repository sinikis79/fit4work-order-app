import csv
import logging
import os
import shutil
import smtplib
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from threading import Lock

import openpyxl
import requests
from dotenv import load_dotenv
from filelock import FileLock
from flask import (
    Flask,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    session,
    url_for,
)
from PIL import Image, ImageDraw, ImageFont, ImageOps
from werkzeug.utils import secure_filename


load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
TEMPLATES_DIR = BASE_DIR / "templates"
STATIC_DIR = BASE_DIR / "static"
CSS_DIR = STATIC_DIR / "css"
JS_DIR = STATIC_DIR / "js"
IMAGE_DIR = STATIC_DIR / "product_images"
BACKUP_DIR = BASE_DIR / "backups"
IMAGE_BACKUP_DIR = BACKUP_DIR / "images"
ORDER_BACKUP_DIR = BACKUP_DIR / "orders"
PRODUCTS_FILE = BASE_DIR / "products.xlsx"
ORDERS_LOG_FILE = BASE_DIR / "orders_log.csv"
FAILED_ORDERS_FILE = BASE_DIR / "failed_orders.csv"
ORDERS_LOG_LOCK_FILE = BASE_DIR / "orders_log.csv.lock"
FAILED_ORDERS_LOCK_FILE = BASE_DIR / "failed_orders.csv.lock"
PLACEHOLDER_FILE = IMAGE_DIR / "placeholder.webp"

PRODUCT_HEADERS = ["품목코드", "상품명", "설명", "카테고리", "도매가", "사이즈", "이미지파일", "노출", "정렬순서"]
PRODUCT_REQUIRED_HEADERS = ["품목코드", "상품명", "카테고리", "도매가", "사이즈", "이미지파일", "노출"]
ORDER_HEADERS = [
    "created_at",
    "order_id",
    "customer_name",
    "items",
    "total_qty",
    "total_amount",
    "memo",
    "teams_status",
    "email_alert_status",
]
FAILED_ORDER_HEADERS = [
    "created_at",
    "order_id",
    "customer_name",
    "items",
    "total_qty",
    "total_amount",
    "memo",
    "error_message",
]
ALLOWED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}

FREE_SHIPPING_THRESHOLD = int(os.getenv("FREE_SHIPPING_THRESHOLD", "100000"))
DEFAULT_PORT = int(os.getenv("APP_PORT", "5001"))

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("fit4work-order-app")
ORDER_LOCK = Lock()
ORDERS_LOG_FILE_LOCK = FileLock(str(ORDERS_LOG_LOCK_FILE))
FAILED_ORDERS_FILE_LOCK = FileLock(str(FAILED_ORDERS_LOCK_FILE))

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "change_this_secret")
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024


def ensure_directories() -> None:
    for path in [TEMPLATES_DIR, STATIC_DIR, CSS_DIR, JS_DIR, IMAGE_DIR, BACKUP_DIR, IMAGE_BACKUP_DIR, ORDER_BACKUP_DIR]:
        path.mkdir(parents=True, exist_ok=True)


def ensure_csv(path: Path, headers: list[str]) -> None:
    if not path.exists():
        with path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file)
            writer.writerow(headers)
        return

    with path.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.reader(file)
        try:
            existing_headers = next(reader)
        except StopIteration:
            existing_headers = []

    if existing_headers == headers:
        return

    rows = read_csv_rows(path) if existing_headers else []
    rows.reverse()
    normalized_rows = [normalize_csv_row(row, headers) for row in rows]
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=headers)
        writer.writeheader()
        writer.writerows(normalized_rows)


def create_sample_products_xlsx() -> None:
    if PRODUCTS_FILE.exists():
        return

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "products"
    sheet.append(PRODUCT_HEADERS)
    sheet.append(["AP-1001-1", "베이직 앞치마", "가볍고 관리가 쉬운 기본 앞치마", "앞치마", 12000, "S,M,L", "AP-1001-1.webp", "Y", 10])
    sheet.append(["CC-2001-2", "조리복 상의", "매장 유니폼으로 쓰기 좋은 조리복", "조리복", 28000, "55,66,77", "CC-2001-2.webp", "Y", 20])
    sheet.append(["HT-3001-1", "위생모", "깔끔한 착용감의 기본 위생모", "모자", 7000, "FREE", "HT-3001-1.webp", "Y", 30])
    workbook.save(PRODUCTS_FILE)


def create_placeholder_image() -> None:
    if PLACEHOLDER_FILE.exists():
        return

    image = Image.new("RGB", (400, 400), "#f1efe7")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((28, 28, 372, 372), radius=32, outline="#d2cbbb", width=4, fill="#faf8f2")
    draw.rectangle((122, 104, 278, 240), outline="#8b9a75", width=6)
    draw.line((122, 240, 200, 170, 278, 240), fill="#8b9a75", width=6)
    text = "Fit4Work"
    subtext = "No Image"
    font = ImageFont.load_default()
    draw.text((145, 280), text, fill="#2d4a30", font=font)
    draw.text((156, 305), subtext, fill="#6b735f", font=font)
    image.save(PLACEHOLDER_FILE, format="WEBP", quality=80, method=6)


def initialize_files() -> None:
    ensure_directories()
    ensure_csv(ORDERS_LOG_FILE, ORDER_HEADERS)
    ensure_csv(FAILED_ORDERS_FILE, FAILED_ORDER_HEADERS)
    backup_csv_files()
    create_sample_products_xlsx()
    create_placeholder_image()


def backup_csv_file(source: Path, prefix: str) -> None:
    if not source.exists():
        return

    backup_path = ORDER_BACKUP_DIR / f"{prefix}_{datetime.now():%Y%m%d}.csv"
    if backup_path.exists():
        return

    shutil.copy2(source, backup_path)


def backup_csv_files() -> None:
    backup_csv_file(ORDERS_LOG_FILE, "orders_log")
    backup_csv_file(FAILED_ORDERS_FILE, "failed_orders")


def safe_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_int(value, default: int = 0) -> int:
    if value in (None, ""):
        return default
    text = str(value).replace(",", "").replace("원", "").strip()
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return default


def format_currency(amount: int) -> str:
    return f"{amount:,}원"


def sort_key_for_product(product: dict) -> tuple:
    order = product.get("sort_order")
    if order is None:
        return (999999, product["product_code"])
    return (order, product["product_code"])


def validate_products_headers(headers: list[str]) -> list[str]:
    normalized = [safe_text(header) for header in headers]
    return [header for header in PRODUCT_REQUIRED_HEADERS if header not in normalized]


def normalize_product_row(row: dict) -> dict | None:
    product_code = safe_text(row.get("품목코드"))
    product_name = safe_text(row.get("상품명"))
    if not product_code or not product_name:
        return None

    exposure = safe_text(row.get("노출")).upper() or "N"
    image_name = safe_text(row.get("이미지파일")) or "placeholder.webp"
    raw_sizes = safe_text(row.get("사이즈"))
    sizes = [size.strip() for size in raw_sizes.split(",") if size and size.strip()]
    if not sizes:
        sizes = ["FREE"]

    sort_order_raw = safe_text(row.get("정렬순서"))
    sort_order = parse_int(sort_order_raw, default=0) if sort_order_raw else None
    unit_price = parse_int(row.get("도매가"))
    image_path = IMAGE_DIR / image_name
    image_exists = image_path.exists()
    image_url = f"/static/product_images/{image_name}" if image_exists else "/static/product_images/placeholder.webp"

    return {
        "product_code": product_code,
        "product_name": product_name,
        "description": safe_text(row.get("설명")),
        "category": safe_text(row.get("카테고리")),
        "unit_price": unit_price,
        "sizes": sizes,
        "image_file": image_name,
        "is_visible": exposure == "Y",
        "sort_order": sort_order,
        "image_exists": image_exists,
        "image_url": image_url,
    }


def load_products() -> list[dict]:
    if not PRODUCTS_FILE.exists():
        return []

    try:
        workbook = openpyxl.load_workbook(PRODUCTS_FILE, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    except Exception:  # noqa: BLE001
        logger.exception("Failed to read products.xlsx")
        return []

    if not rows:
        return []

    headers = [safe_text(cell) for cell in rows[0]]
    products = []
    for row_values in rows[1:]:
        if not row_values or all(value in (None, "") for value in row_values):
            continue
        row = dict(zip(headers, row_values))
        product = normalize_product_row(row)
        if product and product["is_visible"]:
            products.append(product)

    products.sort(key=sort_key_for_product)
    return products


def load_products_for_admin() -> list[dict]:
    if not PRODUCTS_FILE.exists():
        return []

    try:
        workbook = openpyxl.load_workbook(PRODUCTS_FILE, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    except Exception:  # noqa: BLE001
        logger.exception("Failed to read products.xlsx for admin")
        return []

    if not rows:
        return []

    headers = [safe_text(cell) for cell in rows[0]]
    items = []
    for row_values in rows[1:]:
        if not row_values or all(value in (None, "") for value in row_values):
            continue
        row = dict(zip(headers, row_values))
        product = normalize_product_row(row)
        if not product:
            continue
        product["raw_visible"] = safe_text(row.get("노출")).upper() or "N"
        items.append(product)

    items.sort(key=sort_key_for_product)
    return items


def read_csv_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []

    with path.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        rows = list(reader)
    rows.reverse()
    return rows


def normalize_csv_row(row: dict, headers: list[str]) -> dict:
    return {header: row.get(header, "") for header in headers}


def write_csv_rows(path: Path, headers: list[str], rows: list[dict]) -> None:
    normalized_rows = [normalize_csv_row(row, headers) for row in rows]
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=headers)
        writer.writeheader()
        writer.writerows(normalized_rows)


def clear_csv_rows(path: Path, headers: list[str]) -> None:
    write_csv_rows(path, headers, [])


def append_csv_row(path: Path, headers: list[str], row: dict) -> None:
    ensure_csv(path, headers)
    normalized_row = normalize_csv_row(row, headers)
    with path.open("a", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=headers)
        writer.writerow(normalized_row)


def update_order_status(order_id: str, teams_status: str, email_status: str) -> None:
    with ORDERS_LOG_FILE_LOCK:
        rows = read_csv_rows(ORDERS_LOG_FILE)
        rows.reverse()
        updated_rows = []
        for row in rows:
            if row.get("order_id") == order_id:
                row["teams_status"] = teams_status
                row["email_alert_status"] = email_status
            updated_rows.append(row)
        write_csv_rows(ORDERS_LOG_FILE, ORDER_HEADERS, updated_rows)


def generate_order_id() -> str:
    date_key = datetime.now().strftime("%Y%m%d")
    latest_number = 0
    for row in read_csv_rows(ORDERS_LOG_FILE):
        order_id = safe_text(row.get("order_id"))
        if not order_id.startswith(f"F4W-{date_key}-"):
            continue
        try:
            number = int(order_id.split("-")[-1])
        except ValueError:
            continue
        latest_number = max(latest_number, number)
    return f"F4W-{date_key}-{latest_number + 1:03d}"


def build_items_text(items: list[dict]) -> str:
    lines = []
    for item in items:
        lines.append(f"{item['product_code']} / {item['size']} / {item['qty']}장")
    return "\n".join(lines)


def build_teams_payload(order: dict) -> dict:
    lines = [
        "[신규 주문 접수]",
        "",
        f"거래처: {order['customer_name']}",
        f"주문번호: {order['order_id']}",
        "",
        order["items"],
        "",
        f"총수량: {order['total_qty']}장",
        f"총 참고금액: {format_currency(parse_int(order['total_amount']))}",
    ]
    if order.get("memo"):
        lines.extend(["", f"요청사항: {order['memo']}"])
    text = "\n".join(lines)
    return {"text": text}


def send_teams_webhook(order: dict) -> tuple[bool, str]:
    webhook_url = safe_text(os.getenv("TEAMS_WEBHOOK_URL"))
    if not webhook_url:
        return False, "TEAMS_WEBHOOK_URL is empty."

    try:
        response = requests.post(webhook_url, json=build_teams_payload(order), timeout=10)
        response.raise_for_status()
        return True, ""
    except requests.RequestException as exc:
        logger.exception("Teams webhook failed for %s", order["order_id"])
        return False, str(exc)


def send_failure_email(order: dict, error_message: str) -> tuple[bool, str]:
    host = safe_text(os.getenv("EMAIL_HOST"))
    port = parse_int(os.getenv("EMAIL_PORT"), default=587)
    username = safe_text(os.getenv("EMAIL_USER"))
    password = safe_text(os.getenv("EMAIL_PASSWORD"))
    recipient = safe_text(os.getenv("ALERT_EMAIL_TO"))

    if not all([host, port, username, password, recipient]):
        return False, "Email environment values are incomplete."

    message = EmailMessage()
    message["Subject"] = "[Fit4Work] Teams 전송 실패 - 주문 확인 필요"
    message["From"] = username
    message["To"] = recipient
    lines = [
        "Teams 전송에 실패한 주문이 있습니다.",
        "",
        f"주문번호: {order['order_id']}",
        f"거래처: {order['customer_name']}",
        "",
        "주문 내역:",
        order["items"],
        "",
        f"총수량: {order['total_qty']}장",
        f"총 참고금액: {format_currency(parse_int(order['total_amount']))}",
    ]
    if order.get("memo"):
        lines.extend(["", f"요청사항: {order['memo']}"])
    lines.extend(
        [
            "",
            f"오류 내용: {error_message}",
            "",
            "failed_orders.csv와 서버 로그를 확인해주세요.",
        ]
    )
    message.set_content("\n".join(lines))

    try:
        with smtplib.SMTP(host, port, timeout=10) as server:
            server.starttls()
            server.login(username, password)
            server.send_message(message)
        return True, ""
    except Exception as exc:  # noqa: BLE001
        logger.exception("Failure alert email failed for %s", order["order_id"])
        return False, str(exc)


def is_admin_authenticated() -> bool:
    return session.get("admin_authenticated") is True


def normalize_uploaded_image(file_storage) -> str:
    original_name = secure_filename(file_storage.filename or "")
    stem = Path(original_name).stem
    extension = Path(original_name).suffix.lower()
    if extension not in ALLOWED_IMAGE_EXTENSIONS:
        raise ValueError("허용되지 않는 이미지 확장자입니다.")
    if not stem:
        raise ValueError("파일명이 비어 있습니다.")

    target_name = f"{stem}.webp"
    target_path = IMAGE_DIR / target_name
    if target_path.exists():
        backup_name = f"{stem}_{datetime.now():%Y%m%d_%H%M%S}.webp"
        shutil.copy2(target_path, IMAGE_BACKUP_DIR / backup_name)

    image = Image.open(file_storage.stream).convert("RGB")
    image = ImageOps.fit(image, (400, 400), method=Image.Resampling.LANCZOS)
    image.save(target_path, format="WEBP", quality=82, method=6)
    return target_name


def backup_existing_products_file() -> None:
    if not PRODUCTS_FILE.exists():
        return
    backup_name = f"products_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    shutil.copy2(PRODUCTS_FILE, BACKUP_DIR / backup_name)


def save_uploaded_products(file_storage) -> tuple[bool, str]:
    temp_path = BACKUP_DIR / f"_validate_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    file_storage.save(temp_path)
    try:
        workbook = openpyxl.load_workbook(temp_path, data_only=True)
        sheet = workbook.active
        header_row = [safe_text(cell) for cell in next(sheet.iter_rows(values_only=True), [])]
        workbook.close()
    except Exception:  # noqa: BLE001
        temp_path.unlink(missing_ok=True)
        return False, "업로드한 엑셀 파일을 읽을 수 없습니다."

    missing_headers = validate_products_headers(header_row)
    if missing_headers:
        temp_path.unlink(missing_ok=True)
        return False, f"필수 컬럼이 누락되었습니다: {', '.join(missing_headers)}"

    backup_existing_products_file()
    shutil.move(str(temp_path), str(PRODUCTS_FILE))
    return True, "products.xlsx 업로드가 완료되었습니다."


def build_order_from_payload(payload: dict) -> tuple[dict | None, list[str]]:
    customer_name = safe_text(payload.get("customer_name"))
    memo = safe_text(payload.get("memo"))
    items = payload.get("items")
    total_qty = parse_int(payload.get("total_qty"))
    total_amount = parse_int(payload.get("total_amount"))
    errors = []

    if not customer_name:
        errors.append("거래처명을 입력해주세요.")
    if not isinstance(items, list) or not items:
        errors.append("주문 수량을 1개 이상 선택해주세요.")

    normalized_items = []
    available_products = {product["product_code"]: product for product in load_products()}
    computed_total_qty = 0
    computed_total_amount = 0

    if isinstance(items, list):
        for item in items:
            product_code = safe_text(item.get("product_code"))
            size = safe_text(item.get("size"))
            qty = parse_int(item.get("qty"))

            if qty <= 0:
                continue

            product = available_products.get(product_code)
            if not product:
                errors.append(f"{product_code} 상품을 찾을 수 없습니다.")
                continue
            if size not in product["sizes"]:
                errors.append(f"{product_code} 상품에 {size} 사이즈가 없습니다.")
                continue

            normalized_items.append(
                {
                    "product_code": product_code,
                    "product_name": product["product_name"],
                    "size": size,
                    "qty": qty,
                    "unit_price": product["unit_price"],
                    "line_total": product["unit_price"] * qty,
                }
            )
            computed_total_qty += qty
            computed_total_amount += product["unit_price"] * qty

    if not normalized_items:
        errors.append("주문 가능한 항목이 없습니다.")

    if total_qty and total_qty != computed_total_qty:
        errors.append("총수량 정보가 올바르지 않습니다.")
    if total_amount and total_amount != computed_total_amount:
        errors.append("총금액 정보가 올바르지 않습니다.")

    if errors:
        return None, list(dict.fromkeys(errors))

    order = {
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "order_id": generate_order_id(),
        "customer_name": customer_name,
        "items": build_items_text(normalized_items),
        "total_qty": str(computed_total_qty),
        "total_amount": str(computed_total_amount),
        "memo": memo,
        "teams_status": "PENDING",
        "email_alert_status": "PENDING",
        "item_rows": normalized_items,
    }
    return order, []


@app.context_processor
def inject_settings():
    return {"free_shipping_threshold": FREE_SHIPPING_THRESHOLD}


@app.get("/")
def order_page():
    initialize_files()
    products = load_products()
    categories = sorted({product["category"] for product in products if product["category"]})
    return render_template("order.html", products=products, categories=categories)


@app.post("/submit_order")
def submit_order():
    initialize_files()
    payload = request.get_json(silent=True) or {}
    with ORDER_LOCK:
        with ORDERS_LOG_FILE_LOCK:
            order, errors = build_order_from_payload(payload)
            if errors:
                return jsonify({"ok": False, "errors": errors}), 400

            order_row = {key: order[key] for key in ORDER_HEADERS}
            append_csv_row(ORDERS_LOG_FILE, ORDER_HEADERS, order_row)

    teams_ok, teams_error = send_teams_webhook(order)
    if teams_ok:
        update_order_status(order["order_id"], "TEAMS_SUCCESS", "NOT_REQUIRED")
        return jsonify({"ok": True, "order_id": order["order_id"], "message": "주문이 접수되었습니다."})

    update_order_status(order["order_id"], "TEAMS_FAIL", "PENDING")
    failed_row = {
        "created_at": order["created_at"],
        "order_id": order["order_id"],
        "customer_name": order["customer_name"],
        "items": order["items"],
        "total_qty": order["total_qty"],
        "total_amount": order["total_amount"],
        "memo": order["memo"],
        "error_message": teams_error,
    }
    with FAILED_ORDERS_FILE_LOCK:
        append_csv_row(FAILED_ORDERS_FILE, FAILED_ORDER_HEADERS, failed_row)

    email_ok, email_error = send_failure_email(order, teams_error)
    final_email_status = "EMAIL_SENT" if email_ok else "EMAIL_FAIL"
    update_order_status(order["order_id"], "TEAMS_FAIL", final_email_status)
    if not email_ok:
        logger.error("Email alert failed for %s: %s", order["order_id"], email_error)

    return (
        jsonify(
            {
                "ok": True,
                "order_id": order["order_id"],
                "message": "주문은 저장되었고 접수되었습니다. Teams 전송이 실패해 관리자 알림을 처리했습니다.",
                "teams_status": "TEAMS_FAIL",
                "email_alert_status": final_email_status,
            }
        ),
        202,
    )


@app.route("/admin", methods=["GET", "POST"])
def admin_page():
    initialize_files()

    if request.method == "POST" and not is_admin_authenticated():
        admin_password = safe_text(os.getenv("ADMIN_PASSWORD"))
        submitted_password = safe_text(request.form.get("password"))
        if admin_password and submitted_password == admin_password:
            session["admin_authenticated"] = True
            flash("관리자 로그인에 성공했습니다.", "success")
            return redirect(url_for("admin_page"))
        if not admin_password:
            flash("ADMIN_PASSWORD 환경변수를 설정해주세요.", "error")
            return redirect(url_for("admin_page"))
        flash("비밀번호가 올바르지 않습니다.", "error")

    if not is_admin_authenticated():
        return render_template("admin.html", authenticated=False, products=[], orders=[], failed_orders=[], stats={})

    orders = read_csv_rows(ORDERS_LOG_FILE)[:100]
    failed_orders = read_csv_rows(FAILED_ORDERS_FILE)[:100]
    for failed_order in failed_orders:
        failed_order["teams_status"] = "TEAMS_FAIL"
    products = load_products_for_admin()
    stats = {
        "product_count": len(products),
        "order_count": len(read_csv_rows(ORDERS_LOG_FILE)),
        "failed_count": len(read_csv_rows(FAILED_ORDERS_FILE)),
        "sales_total": sum(parse_int(row.get("total_amount")) for row in read_csv_rows(ORDERS_LOG_FILE)),
    }
    return render_template(
        "admin.html",
        authenticated=True,
        products=products,
        orders=orders,
        failed_orders=failed_orders,
        stats=stats,
    )


@app.post("/admin/upload-products")
def upload_products():
    if not is_admin_authenticated():
        return redirect(url_for("admin_page"))

    file_storage = request.files.get("products_file")
    if not file_storage or not safe_text(file_storage.filename):
        flash("업로드할 products.xlsx 파일을 선택해주세요.", "error")
        return redirect(url_for("admin_page"))

    ok, message = save_uploaded_products(file_storage)
    flash(message, "success" if ok else "error")
    return redirect(url_for("admin_page"))


@app.post("/admin/upload-images")
def upload_images():
    if not is_admin_authenticated():
        return redirect(url_for("admin_page"))

    files = request.files.getlist("image_files")
    if not files or not any(safe_text(file.filename) for file in files):
        flash("업로드할 이미지 파일을 선택해주세요.", "error")
        return redirect(url_for("admin_page"))

    saved_files = []
    errors = []
    for file_storage in files:
        if not safe_text(file_storage.filename):
            continue
        try:
            saved_files.append(normalize_uploaded_image(file_storage))
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{file_storage.filename}: {exc}")

    if saved_files:
        flash(f"이미지 {len(saved_files)}건 업로드 완료: {', '.join(saved_files)}", "success")
    if errors:
        flash(" / ".join(errors), "error")
    return redirect(url_for("admin_page"))


@app.post("/admin/clear-failed-orders")
def clear_failed_orders():
    if not is_admin_authenticated():
        return redirect(url_for("admin_page"))

    with FAILED_ORDERS_FILE_LOCK:
        clear_csv_rows(FAILED_ORDERS_FILE, FAILED_ORDER_HEADERS)
    flash("실패 주문 로그를 초기화했습니다.", "success")
    return redirect(url_for("admin_page"))


@app.post("/admin/clear-orders")
def clear_orders():
    if not is_admin_authenticated():
        return redirect(url_for("admin_page"))

    with ORDERS_LOG_FILE_LOCK:
        clear_csv_rows(ORDERS_LOG_FILE, ORDER_HEADERS)
    flash("전체 주문 로그를 초기화했습니다.", "success")
    return redirect(url_for("admin_page"))


@app.get("/admin/logout")
def admin_logout():
    session.pop("admin_authenticated", None)
    flash("관리자 세션이 종료되었습니다.", "success")
    return redirect(url_for("admin_page"))


@app.errorhandler(404)
def not_found(_error):
    return jsonify({"ok": False, "error": "Not found"}), 404


initialize_files()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=DEFAULT_PORT, debug=os.getenv("FLASK_ENV") == "development")
